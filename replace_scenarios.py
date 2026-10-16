"""Run from ~/SA_APIs:  python3 replace_scenarios.py portfolio_aware_scenarios.xlsx
Replaces the 'Portfolio aware recommendations' sheet of the framework test data workbook with the
new one (150 rows incl. monthly income, with CASE ID). Backup is saved first."""
import sys, shutil
from datetime import datetime
import openpyxl
TARGET = "src/test/resources/testdata/Portfolio Planner Test CasesQA .xlsx"
SHEET = "Portfolio aware recommendations"
if len(sys.argv) < 2:
    sys.exit(__doc__)
src = openpyxl.load_workbook(sys.argv[1], data_only=True)
sws = src[SHEET] if SHEET in src.sheetnames else src.worksheets[0]
wb = openpyxl.load_workbook(TARGET)
pos = wb.sheetnames.index(SHEET) if SHEET in wb.sheetnames else len(wb.sheetnames)
if SHEET in wb.sheetnames:
    del wb[SHEET]
ws = wb.create_sheet(SHEET, pos)
for row in sws.iter_rows(values_only=True):
    ws.append(list(row))
backup = TARGET.replace(".xlsx", f".backup-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx")
shutil.copy(TARGET, backup)
wb.save(TARGET)
print(f"Replaced '{SHEET}' with {sws.max_row - 1} rows. Backup: {backup}")
