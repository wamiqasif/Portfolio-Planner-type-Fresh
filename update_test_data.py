"""
Update the framework's test data from a new sheet sent by the development team.

Usage (from ~/SA_APIs):
    python3 update_test_data.py <new_data_file>            # apply
    python3 update_test_data.py <new_data_file> --dry-run  # only show what would change
    python3 update_test_data.py <dev_file.xlsx> --sheet "Existing Account Revised List" --sheet "Regular Income Accounts - Ashut"
                                                           # read specific tabs of the dev's workbook

<new_data_file> can be .xlsx (first sheet, or a sheet named "Test Accounts"), .csv, or .tsv/.txt
(copy the rows from Excel / chat and paste them into a text file - tab separated, header row first).

What it does:
  * Opens src/test/resources/testdata/Portfolio Planner Test CasesQA .xlsx, sheet "Test Accounts".
  * For every "Case ID" in the new file, overwrites that row's cells for every column whose header
    exists in both files (Investor ID, User ID, Label ID, Plan ID, Matched Fund, Opinion Name,
    Star Rating, Is Analyst Pick, Is Direct Growth, Login URL, ...). Other rows are not touched.
  * Columns that only exist in the new file (e.g. "Remark", "Label Name") are added at the end;
    the framework reads by header name, so extra columns are harmless.
  * Case IDs not present yet are appended as new rows.
  * Regular Income Case IDs "RI-1".."RI-17" are mapped to 201..217 automatically.
  * Saves a timestamped backup first and prints a change log (also written to update_test_data.log).
"""
import csv
import re
import os
import shutil
import sys
from datetime import datetime

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl missing. Install with:  pip3 install --user openpyxl")

TARGET = "src/test/resources/testdata/Portfolio Planner Test CasesQA .xlsx"
SHEET = "Test Accounts"
KEY = "Case ID"


def norm(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def read_new(path, sheet=None):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, data_only=True)
        if sheet:
            if sheet not in wb.sheetnames:
                sys.exit(f"Sheet '{sheet}' not found. Sheets in file: {wb.sheetnames}")
            ws = wb[sheet]
        else:
            ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        rows = [[norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, encoding="utf-8-sig", newline="") as f:
            text = f.read()
        delim = "\t" if text.count("\t") >= text.count(",") else ","
        rows = [[norm(c) for c in r] for r in csv.reader(text.splitlines(), delimiter=delim)]
    rows = [r for r in rows if any(r)]
    if not rows:
        sys.exit("New data file is empty.")
    header = rows[0]
    if KEY not in header:
        sys.exit(f"New data file has no '{KEY}' column. Header found: {header}")
    out = []
    for r in rows[1:]:
        rec = {h: (r[i] if i < len(r) else "") for i, h in enumerate(header) if h}
        cid = rec.get(KEY, "")
        m = re.match(r"^RI-(\d+)$", cid, re.IGNORECASE)
        if m:  # Regular Income sheet uses RI-1..RI-17; framework uses numeric Case IDs 201..217
            rec["Source Case"] = cid
            rec[KEY] = str(200 + int(m.group(1)))
        if rec.get(KEY):
            out.append(rec)
    if any("Source Case" in r for r in out) and "Source Case" not in header:
        header = header + ["Source Case"]
    return header, out


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    dry = "--dry-run" in sys.argv
    sheets = [sys.argv[i + 1] for i, a in enumerate(sys.argv) if a == "--sheet" and i + 1 < len(sys.argv)]
    if not os.path.exists(TARGET):
        sys.exit(f"Target not found: {TARGET} (run this from the SA_APIs folder)")

    new_header, new_rows = [], []
    for sh in (sheets or [None]):
        h, rows = read_new(src, sh)
        new_header += [x for x in h if x not in new_header]
        new_rows += rows
        if sh:
            print(f"Read {len(rows)} rows from sheet '{sh}'")
    wb = openpyxl.load_workbook(TARGET)
    ws = wb[SHEET]
    header = [norm(c.value) for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header) if h}

    for h in new_header:  # add columns that exist only in the new file
        if h and h not in col:
            c = ws.max_column + 1
            ws.cell(1, c, h)
            col[h] = c

    row_of = {}
    for r in range(2, ws.max_row + 1):
        k = norm(ws.cell(r, col[KEY]).value)
        if k:
            row_of[k] = r

    log = []
    updated = appended = 0
    for rec in new_rows:
        cid = rec[KEY]
        r = row_of.get(cid)
        if r is None:
            r = ws.max_row + 1
            row_of[cid] = r
            appended += 1
            log.append(f"Case {cid}: NEW ROW")
        changes = []
        for h, v in rec.items():
            if not h:
                continue
            old = norm(ws.cell(r, col[h]).value)
            if old != v:
                changes.append(f"{h}: '{old}' -> '{v}'")
                if not dry:
                    ws.cell(r, col[h], v if v != "" else None)
        if changes:
            updated += 1
            log.append(f"Case {cid}: " + "; ".join(changes))

    print("\n".join(log) if log else "No changes.")
    print(f"\n{len(new_rows)} rows in new file | {updated} cases changed | {appended} new cases")
    if dry:
        print("DRY RUN - nothing saved.")
        return
    backup = TARGET.replace(".xlsx", f".backup-{datetime.now():%Y%m%d-%H%M%S-%f}.xlsx")
    shutil.copy(TARGET, backup)
    wb.save(TARGET)
    with open("update_test_data.log", "a", encoding="utf-8") as f:
        f.write(f"\n=== {datetime.now():%Y-%m-%d %H:%M} from {src}\n" + "\n".join(log) + "\n")
    print(f"Saved. Backup: {backup}\nLog: update_test_data.log")


if __name__ == "__main__":
    main()
