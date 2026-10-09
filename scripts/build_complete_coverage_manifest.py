#!/usr/bin/env python3
"""
Builds src/test/resources/testdata/generated/complete-coverage-manifest.json: the execution mapping for
the complete Portfolio Planner business-rule coverage run. READ-ONLY on every source workbook.

Sources
  1. ~/Downloads/Portfolio Planner test cases(1).xlsx
       - "Migration to new code"            (value lists per column, NOT one-test-per-row; see below)
       - "Portfolio aware recommendations"   (133 rows, "Investor account" column empty in every row)
  2. ~/Downloads/Portfolio_Planner_API1_Portfolio_Aware_Test_Data.xlsx   (API-1 boundary/negative data)
  3. src/test/resources/testdata/Portfolio Planner Test CasesQA .xlsx, sheet "Test Accounts"
       - the repo's live account list, used ONLY to supply label/user/investor/login-url for the
         Portfolio-Aware rows (the new workbook has no account for any row).

Usage: python3 scripts/build_complete_coverage_manifest.py [--income-map path.json]
"""
import json, re, sys, collections, os
import openpyxl

HOME = os.path.expanduser("~")
TESTCASES = f"{HOME}/Downloads/Portfolio Planner test cases(1).xlsx"
API1 = f"{HOME}/Downloads/Portfolio_Planner_API1_Portfolio_Aware_Test_Data.xlsx"
REPO_WB = "src/test/resources/testdata/Portfolio Planner Test CasesQA .xlsx"
OUT = "src/test/resources/testdata/generated/complete-coverage-manifest.json"

STATIC = {"label_id": 89660240, "user_id": 123, "investor_id": 87683240}
BASE = {"duration_type": "y", "investment_duration": 15, "risk_profile_id": 2, "annual_income_range": 2}
NEEDED_ANNUAL_DEFAULT = 5000  # mandatory for REGULAR-INCOME (422 otherwise) - not in the API-1 workbook


def norm(s):
    if s is None:
        return ""
    if isinstance(s, (int, float)) and float(s) == int(s):
        return str(int(s))
    return re.sub(r"\s+", " ", str(s)).strip()


def num(v):
    if v is None or str(v).strip().upper() in ("", "NA"):
        return None
    return int(float(v))


def payload(goal, txn=None, dur_type="y", dur=15, monthly=None, lumpsum=None, accumulated=None, tax=None,
            needed=None, risk=2, income=2, ident=None):
    ident = ident or STATIC
    p = {"goal_type": goal}
    if txn:
        p["txn_options"] = txn
    p["duration_type"] = dur_type
    p["investment_duration"] = dur
    if monthly is not None:
        p["monthly_amount"] = monthly
    if lumpsum is not None:
        p["lumpsum_amount"] = lumpsum
    if accumulated is not None:
        p["accumulated_amount"] = accumulated
    if needed is not None:
        p["needed_annual_amount"] = needed
    if tax is not None:
        p["tax_saving_amount"] = tax
    p["label_id"] = ident["label_id"]
    p["user_id"] = ident["user_id"]
    p["risk_profile_id"] = risk
    p["annual_income_range"] = income
    p["aware_type"] = "PORTFOLIO_AWARE"
    p["investor_id"] = ident["investor_id"]
    return p


def expected_valid_by_api1_rules(p):
    """Independent restatement of the API-1 Business_Rules sheet. Returns (valid, [reasons])."""
    why = []
    g = p.get("goal_type")
    t = p.get("txn_options")
    if g != "GROWTH" and "txn_options" in p:
        why.append("txn_options only for GROWTH")
    if g == "GROWTH":
        if t not in ("SIP", "LUMPSUM", "SIPLUMPSUM"):
            why.append("GROWTH needs txn_options")
        if t in ("SIP", "SIPLUMPSUM") and "monthly_amount" not in p:
            why.append("monthly missing")
        if t == "LUMPSUM" and "monthly_amount" in p:
            why.append("monthly not allowed with LUMPSUM")
        if t == "SIP" and "lumpsum_amount" in p:
            why.append("lumpsum not allowed with SIP")
        if t in ("LUMPSUM", "SIPLUMPSUM") and "lumpsum_amount" not in p:
            why.append("lumpsum missing")
    else:
        for f in ("monthly_amount", "lumpsum_amount"):
            if f in p:
                why.append(f + " not allowed for non-GROWTH")
    if g != "GROWTH" and "txn_options" not in p:
        pass
    if "accumulated_amount" in p and g != "REGULAR-INCOME":
        why.append("accumulated only for REGULAR-INCOME")
    if g == "REGULAR-INCOME" and "accumulated_amount" not in p:
        why.append("accumulated missing")
    if "tax_saving_amount" in p and g != "TAX-SAVINGS":
        why.append("tax only for TAX-SAVINGS")
    if g == "TAX-SAVINGS" and "tax_saving_amount" not in p:
        why.append("tax missing")
    d, dt = p.get("investment_duration"), p.get("duration_type")
    if dt == "y" and not (1 <= d <= 30):
        why.append(f"duration y {d} outside 1-30")
    if dt == "m" and not (6 <= d <= 359):
        why.append(f"duration m {d} outside 6-359")
    rng = {"monthly_amount": (5000, 500000000), "lumpsum_amount": (1000, 100000000),
           "accumulated_amount": (100000, 100000000), "tax_saving_amount": (500, 150000)}
    for f, (lo, hi) in rng.items():
        if f in p and not (lo <= p[f] <= hi):
            why.append(f"{f} {p[f]} outside {lo}-{hi}")
    return (not why), why


def main():
    income_map = {}
    if "--income-map" in sys.argv:
        income_map = json.load(open(sys.argv[sys.argv.index("--income-map") + 1]))  # {"100000": 1, ...}

    scenarios = []

    def add(**kw):
        kw.setdefault("mappingStatus", "EXECUTABLE")
        kw.setdefault("mappingReason", "")
        kw["scenarioId"] = f"S{len(scenarios) + 1:04d}"
        scenarios.append(kw)

    tc = openpyxl.load_workbook(TESTCASES, data_only=True)

    # ------------------------------------------------------------------ Portfolio-Aware sheet (133)
    repo = openpyxl.load_workbook(REPO_WB, data_only=True)
    accounts = [r for r in list(repo["Test Accounts"].iter_rows(values_only=True))[1:] if r[0]]
    hdr = [c.value for c in repo["Test Accounts"][1]]
    ix = {h: i for i, h in enumerate(hdr) if h}
    ta_group = {"Flexicap/Multicap/Large & Mid cap/Focused": "Equity: Flexi Cap", "Short Duration": "Debt: Short Term"}
    used = set()
    pa_rows = list(tc["Portfolio aware recommendations"].iter_rows(values_only=True))
    for ridx, r in enumerate(pa_rows[1:], start=2):
        cat, case, goal, risk, years, lump, mon, tax = (norm(r[1]), norm(r[2]), norm(r[3]), norm(r[4]),
                                                          r[5], r[6], r[7], r[8])
        want = ta_group.get(cat, cat)
        acct = None
        for i, a in enumerate(accounts):
            if i in used:
                continue
            if norm(a[ix["Fund Category"]]) == want and norm(a[ix["Cases"]]) == case:
                used.add(i)
                acct = a
                break
        base = dict(sourceSheet="Portfolio aware recommendations", sourceRow=ridx, tcId=f"PA-{ridx - 1:03d}",
                    fundCategory=cat, caseCovered=case, goal=goal, riskProfile=risk,
                    timeText=norm(years), lumpsumText=norm(lump), monthlyText=norm(mon), taxText=norm(tax),
                    auth="GUID", kind="PORTFOLIO_AWARE_ROW")
        if acct is None:
            add(**base, investorAccount=None, payload=None, expectedValid=None, mappingStatus="DATA_MAPPING_REQUIRED",
                mappingReason=f"No account in 'Test Accounts' for category '{cat}' + case '{case}'")
            continue
        account = {k: norm(acct[ix[k]]) for k in ("Case ID", "Match Status", "Investor ID", "User ID", "Label ID",
                                                  "Plan ID", "Matched Fund", "Is Analyst Pick", "Is Direct Growth",
                                                  "Star Rating", "Login URL", "Sum Units Flow", "Fund Category")}
        if not (account["Label ID"] and account["User ID"] and account["Investor ID"] and account["Login URL"]):
            add(**base, investorAccount=account, payload=None, expectedValid=None, mappingStatus="DATA_MAPPING_REQUIRED",
                mappingReason=f"Test Accounts case {account['Case ID']} is {account['Match Status']} with no label/user/investor/login-url - no account exists for this case (not fabricated)")
            continue
        ident = {"label_id": int(account["Label ID"]), "user_id": int(account["User ID"]),
                 "investor_id": int(account["Investor ID"])}
        risk_id = {"conservative": 1, "moderate": 2, "aggressive": 3}.get(risk.lower())
        notes = []
        if risk_id is None:
            add(**base, investorAccount=account, payload=None, expectedValid=None,
                mappingStatus="DATA_MAPPING_REQUIRED", mappingReason=f"Unknown risk profile '{risk}'")
            continue
        dur = num(years)
        if dur is None:
            dur = BASE["investment_duration"]
            notes.append("Time=NA in Excel: API-1 baseline duration y/15 applied")
        notes.append(f"Excel has no annual-income column: API-1 baseline annual_income_range={BASE['annual_income_range']} applied")
        if goal.lower() == "invest to save tax":
            p = payload("TAX-SAVINGS", dur_type="y", dur=dur, tax=num(tax), risk=risk_id, income=BASE["annual_income_range"], ident=ident)
        else:
            m, l = num(mon), num(lump)
            m = m if m else None  # 0 => field must be ABSENT (never send 0)
            l = l if l else None
            txn = "SIPLUMPSUM" if (m and l) else ("SIP" if m else "LUMPSUM")
            p = payload("GROWTH", txn, "y", dur, monthly=m, lumpsum=l, risk=risk_id, income=BASE["annual_income_range"], ident=ident)
        ev, why = expected_valid_by_api1_rules(p)
        add(**base, investorAccount=account, payload=p, expectedValid=ev, expectedInvalidWhy="; ".join(why),
            notes="; ".join(notes))

    # ------------------------------------------------------------------ Migration sheet
    ws = tc["Migration to new code"]
    rows = list(ws.iter_rows(values_only=True))
    section = None
    SECTION_TITLES = {"Investing for higher returns": "GROWTH", "Investing for monthly income": "REGULAR-INCOME",
                      "Invest to save tax": "TAX-SAVINGS"}
    risk_ids = {"Conservative": 1, "Moderate": 2, "Aggressive": 3}

    def mig(ridx, **kw):
        add(sourceSheet="Migration to new code", sourceRow=ridx, tcId=f"MIG-{ridx:02d}", investorAccount=None,
            fundCategory="", auth="VRLOGIN", kind="MIGRATION_CELL", **kw)

    for ridx, r in enumerate(rows, start=1):
        vals = [norm(v) for v in r]
        if not any(vals):
            add(sourceSheet="Migration to new code", sourceRow=ridx, tcId=f"MIG-{ridx:02d}", investorAccount=None,
                fundCategory="", caseCovered="(blank separator row)", kind="NOT_A_TEST_ROW", payload=None,
                expectedValid=None, mappingStatus="NOT_A_TEST_ROW", mappingReason="Blank separator row")
            continue
        if vals[0] in SECTION_TITLES and not any(vals[1:]):
            section = SECTION_TITLES[vals[0]]
            add(sourceSheet="Migration to new code", sourceRow=ridx, tcId=f"MIG-{ridx:02d}", investorAccount=None,
                fundCategory="", caseCovered=f"(section title: {vals[0]})", kind="NOT_A_TEST_ROW", payload=None,
                expectedValid=None, mappingStatus="NOT_A_TEST_ROW", mappingReason="Section title row")
            continue
        if vals[0] in ("Risk profile",):
            add(sourceSheet="Migration to new code", sourceRow=ridx, tcId=f"MIG-{ridx:02d}", investorAccount=None,
                fundCategory="", caseCovered="(column header row: " + " | ".join(v for v in vals if v) + ")",
                kind="NOT_A_TEST_ROW", payload=None, expectedValid=None, mappingStatus="NOT_A_TEST_ROW",
                mappingReason="Column header row")
            continue
        # a data row: each non-blank cell is one input value to test, one-factor-at-a-time vs the API-1 baseline.
        cells = []  # (label, scenario-or-None, reason)
        if section == "GROWTH":
            risk, income, lump, sip, months = r[0], r[1], r[2], r[3], r[4]
            if risk:
                cells.append(("Risk profile=" + risk, payload("GROWTH", "SIP", monthly=5000, risk=risk_ids[risk]), None))
            if income is not None:
                key = str(int(income))
                if key in income_map:
                    cells.append((f"Annual income={income}", payload("GROWTH", "SIP", monthly=5000, income=income_map[key]), None))
                else:
                    cells.append((f"Annual income={income}", None,
                                  "Rupee annual-income -> annual_income_range code (1-5) mapping is not defined in any "
                                  "source workbook/PRD table (6 rupee values vs 5 codes); Config Master mapping not resolved"))
            if lump is not None:
                if int(lump) == 0:
                    cells.append(("Lumpsum=0 (not selected)", payload("GROWTH", "SIP", monthly=5000),
                                  "NOTE: 0 means 'not chosen' -> lumpsum_amount ABSENT (SIP-only), never sent as 0"))
                else:
                    cells.append((f"Lumpsum={int(lump)}", payload("GROWTH", "LUMPSUM", lumpsum=int(lump)), None))
            if sip is not None:
                if int(sip) == 0:
                    cells.append(("SIP=0 (not selected)", payload("GROWTH", "LUMPSUM", lumpsum=50000),
                                  "NOTE: 0 means 'not chosen' -> monthly_amount ABSENT (LUMPSUM-only), never sent as 0"))
                else:
                    cells.append((f"SIP={int(sip)}", payload("GROWTH", "SIP", monthly=int(sip)), None))
            if months is not None:
                cells.append((f"No. of months={int(months)}", payload("GROWTH", "SIP", "m", int(months), monthly=5000), None))
        elif section == "REGULAR-INCOME":
            risk, when, acc, inc = r[0], r[1], r[2], r[3]
            if risk:
                cells.append(("Risk profile=" + risk, payload("REGULAR-INCOME", accumulated=100000, needed=NEEDED_ANNUAL_DEFAULT, risk=risk_ids[risk]), None))
            if when:
                cells.append((f"Needed when={when}", None,
                              "No Planner Generate request field corresponds to 'Needed when' (no timing/withdrawal-start parameter in the API-1 contract)"))
            if acc is not None:
                cells.append((f"Accumulation={int(acc)}", payload("REGULAR-INCOME", accumulated=int(acc), needed=NEEDED_ANNUAL_DEFAULT), None))
            if inc is not None:
                cells.append((f"Monthly income needed={int(inc)}",
                              payload("REGULAR-INCOME", accumulated=100000, needed=int(inc)),
                              "ASSUMPTION: sent as needed_annual_amount unchanged (documented UI range 500-500000 matches monthly-income magnitudes; field name says 'annual')"))
        elif section == "TAX-SAVINGS":
            risk, month, fy, amt = r[0], r[1], r[2], r[3]
            if risk:
                cells.append(("Risk profile=" + risk, payload("TAX-SAVINGS", tax=500, risk=risk_ids[risk]), None))
            if month:
                cells.append((f"Month={month}", None,
                              "Planner Generate has no month/date parameter; month-dependent behaviour is server-clock driven and cannot be set by the request"))
            if fy:
                cells.append((f"Financial year={fy}", None,
                              "Planner Generate has no financial_year parameter (PRD: UI-only question shown 1 Jan-31 Mar)"))
            if amt is not None:
                cells.append((f"Tax saving amount={int(amt)}", payload("TAX-SAVINGS", tax=int(amt)), None))
        for label, p, note in cells:
            if p is None:
                mig(ridx, caseCovered=label, goal=section, riskProfile="", timeText="", lumpsumText="", monthlyText="",
                    taxText="", payload=None, expectedValid=None, mappingStatus="DATA_MAPPING_REQUIRED", mappingReason=note)
            else:
                ev, why = expected_valid_by_api1_rules(p)
                mig(ridx, caseCovered=label, goal=section,
                    riskProfile=str(p["risk_profile_id"]), timeText=f'{p["investment_duration"]}{p["duration_type"]}',
                    lumpsumText=str(p.get("lumpsum_amount", "")), monthlyText=str(p.get("monthly_amount", "")),
                    taxText=str(p.get("tax_saving_amount", "")), payload=p, expectedValid=ev,
                    expectedInvalidWhy="; ".join(why), notes=(note or ""))

    # ------------------------------------------------------------------ API-1 workbook
    a1 = openpyxl.load_workbook(API1, data_only=True)["API1_Test_Data"]
    rows = list(a1.iter_rows(values_only=True))
    h = [c for c in rows[0]]
    for ridx, r in enumerate(rows[1:], start=2):
        d = dict(zip(h, r))
        p = {"goal_type": d["goal_type"]}
        if d["txn_options"]:
            p["txn_options"] = d["txn_options"]
        p["duration_type"] = d["duration_type"]
        p["investment_duration"] = d["investment_duration"]
        for f in ("monthly_amount", "lumpsum_amount", "accumulated_amount", "tax_saving_amount"):
            if d[f] is not None:
                p[f] = int(d[f])
        notes = []
        if d["goal_type"] == "REGULAR-INCOME":
            p["needed_annual_amount"] = NEEDED_ANNUAL_DEFAULT
            notes.append(f"needed_annual_amount={NEEDED_ANNUAL_DEFAULT} added: mandatory for REGULAR-INCOME (HTTP 422 otherwise), absent from API-1 workbook")
        p["label_id"] = int(d["label_id"]); p["user_id"] = int(d["user_id"])
        p["risk_profile_id"] = int(d["risk_profile_id"]); p["annual_income_range"] = int(d["annual_income_range"])
        p["aware_type"] = d["aware_type"]; p["investor_id"] = int(d["investor_id"])
        ev_wb = d["EXPECTED_RESULT"] == "VALID"
        ev_rules, why = expected_valid_by_api1_rules(p)
        if ev_wb != ev_rules:
            notes.append(f"WORKBOOK/RULES DISAGREE: workbook says {d['EXPECTED_RESULT']}, Business_Rules sheet implies {'VALID' if ev_rules else 'INVALID: ' + '; '.join(why)}")
        add(sourceSheet="API1_Test_Data", sourceRow=ridx, tcId=d["TC_ID"], investorAccount=None, fundCategory="",
            caseCovered=f'{d["CASE_TYPE"]}: {d["BUSINESS_RULE_TO_VALIDATE"]}', goal=d["goal_type"],
            riskProfile=str(d["risk_profile_id"]), timeText=f'{d["investment_duration"]}{d["duration_type"]}',
            lumpsumText=norm(d["lumpsum_amount"]), monthlyText=norm(d["monthly_amount"]), taxText=norm(d["tax_saving_amount"]),
            auth="VRLOGIN", kind="API1_ROW", payload=p, expectedValid=ev_wb, expectedInvalidWhy="; ".join(why),
            notes="; ".join(notes), workbookCaseType=d["CASE_TYPE"])

    # supplementary BVA rows: the four "+1" boundaries the brief lists but the API-1 workbook does not contain
    supp = [("S01", "monthly_amount", 5001, payload("GROWTH", "SIP", monthly=5001)),
            ("S02", "lumpsum_amount", 1001, payload("GROWTH", "LUMPSUM", lumpsum=1001)),
            ("S03", "accumulated_amount", 100001, payload("REGULAR-INCOME", accumulated=100001, needed=NEEDED_ANNUAL_DEFAULT)),
            ("S04", "tax_saving_amount", 501, payload("TAX-SAVINGS", tax=501))]
    for i, (sid, field, val, p) in enumerate(supp):
        ev, why = expected_valid_by_api1_rules(p)
        add(sourceSheet="API1_Test_Data (supplementary - generated by this run)", sourceRow=0, tcId=f"PP-A1-{sid}",
            investorAccount=None, fundCategory="", caseCovered=f"BOUNDARY_VALID: {field}={val} (brief-required, missing from workbook)",
            goal=p["goal_type"], riskProfile="2", timeText="15y", lumpsumText=str(p.get("lumpsum_amount", "")),
            monthlyText=str(p.get("monthly_amount", "")), taxText=str(p.get("tax_saving_amount", "")), auth="VRLOGIN",
            kind="API1_SUPPLEMENTARY", payload=p, expectedValid=ev, expectedInvalidWhy="; ".join(why),
            notes="Not in the supplied workbook; generated so every boundary listed in the brief is executed",
            workbookCaseType="BOUNDARY_VALID")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump({"scenarios": scenarios, "unusedTestAccounts": [norm(accounts[i][0]) for i in range(len(accounts)) if i not in used]},
              open(OUT, "w"), indent=1)

    # ------------------------------------------------------------------ coverage matrix
    c = collections.Counter
    by_sheet = c(s["sourceSheet"] for s in scenarios)
    src_rows = c()
    for s in scenarios:
        src_rows[(s["sourceSheet"], s["sourceRow"])] += 1
    print("SCENARIO LINES (one per executable/blocked input) by sheet:", dict(by_sheet))
    print("Distinct source rows:", c(k[0] for k in src_rows))
    print("Status:", c(s["mappingStatus"] for s in scenarios))
    print("Status by sheet:", {k: dict(c(s["mappingStatus"] for s in scenarios if s["sourceSheet"] == k)) for k in by_sheet})
    print("Unused Test Accounts (not referenced by the 133 rows):", [norm(accounts[i][0]) for i in range(len(accounts)) if i not in used])


if __name__ == "__main__":
    main()
